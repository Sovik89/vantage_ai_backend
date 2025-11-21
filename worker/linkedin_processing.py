"""
worker/linkedin_processing.py
LinkedIn Scout Excel Export Module
Similar to ATS processing but for LinkedIn candidate analysis
"""

from typing import List, Dict, Any
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_linkedin_analysis_to_excel(
    results: List[Dict[str, Any]], 
    output_file: str,
    job_title: str = "Position",
    location: str = "Unknown Location",
    job_description: str = ""
) -> None:
    """
    Export LinkedIn analysis to Excel with multiple sheets
    
    Args:
        results: List of analyzed candidate results
        output_file: Path to save the Excel file
        job_title: Position title
        location: Search location
        job_description: JD text
    """
    
    if not results:
        print("⚠️ No results to export")
        return
    
    wb = Workbook()
    
    # Remove default sheet
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
    
    print("   📄 Creating Summary page...")
    create_linkedin_summary_page(wb, results, job_title, location, job_description)
    
    print("   📄 Creating Candidate Rankings...")
    create_linkedin_rankings_sheet(wb, results)
    
    print("   📄 Creating Skills Analysis...")
    create_linkedin_skills_sheet(wb, results)
    
    print("   📄 Creating Job Hopping Analysis...")
    create_job_hopping_sheet(wb, results)
    
    print("   📄 Creating Experience Details...")
    create_experience_details_sheet(wb, results)
    
    wb.save(output_file)
    print(f"✅ LinkedIn analysis exported to: {output_file}")
    print(f"\n📊 Excel file contains 5 sheets:")
    print(f"   1. Summary - Comprehensive analysis overview")
    print(f"   2. Candidate Rankings - Complete scores")
    print(f"   3. Skills Analysis - Skills matching")
    print(f"   4. Job Hopping Analysis - Stability assessment")
    print(f"   5. Experience Details - Work history")


def create_linkedin_summary_page(
    wb: Workbook, 
    results: List[Dict[str, Any]], 
    job_title: str,
    location: str,
    job_description: str
) -> None:
    """Create comprehensive summary page"""
    
    ws = wb.create_sheet("Summary", 0)
    
    ws.column_dimensions['A'].width = 120
    
    # Header
    ws['A1'] = "LINKEDIN TALENT SCOUT REPORT"
    ws['A1'].font = Font(bold=True, size=16, color="FFFFFF")
    ws['A1'].fill = PatternFill(start_color="0077B5", end_color="0077B5", fill_type="solid")
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells('A1:A2')
    
    row = 3
    
    ws[f'A{row}'] = f"Position: {job_title}"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    row += 1
    
    ws[f'A{row}'] = f"Location: {location}"
    ws[f'A{row}'].font = Font(bold=True, size=12)
    row += 1
    
    ws[f'A{row}'] = f"Analysis Date: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
    ws[f'A{row}'].font = Font(italic=True, size=10)
    row += 2
    
    # Statistics
    ws[f'A{row}'] = "📊 QUICK STATISTICS"
    ws[f'A{row}'].font = Font(bold=True, size=14, color="0077B5")
    row += 1
    
    total = len(results)
    avg_score = sum(r.get('match_score', 0) for r in results) / total if total > 0 else 0
    highly_rec = len([r for r in results if r.get('recommendation') == 'Highly Recommended'])
    rec = len([r for r in results if r.get('recommendation') == 'Recommended'])
    consider = len([r for r in results if r.get('recommendation') == 'Consider'])
    
    stats = [
        ("Total Candidates Analyzed", total),
        ("Average Match Score", f"{avg_score:.1f}%"),
        ("Highly Recommended", highly_rec),
        ("Recommended", rec),
        ("Consider", consider),
        ("Top Match Score", f"{max((r.get('match_score', 0) for r in results), default=0):.1f}%"),
    ]
    
    for label, value in stats:
        ws[f'A{row}'] = f"  • {label}: {value}"
        ws[f'A{row}'].font = Font(size=11)
        row += 1
    
    row += 1
    
    # Scoring methodology
    ws[f'A{row}'] = "⚖️ SCORING METHODOLOGY"
    ws[f'A{row}'].font = Font(bold=True, size=14, color="0077B5")
    row += 1
    
    methodology = [
        "Skills Match - Alignment between candidate skills and job requirements",
        "Experience Score - Years and relevance of experience",
        "Education Score - Educational background and qualifications",
        "Overall Match - Weighted combination of all factors",
    ]
    
    for item in methodology:
        ws[f'A{row}'] = f"  • {item}"
        ws[f'A{row}'].font = Font(size=10)
        row += 1
    
    row += 2
    
    # Top candidates
    ws[f'A{row}'] = "🏆 TOP 5 CANDIDATES AT A GLANCE"
    ws[f'A{row}'].font = Font(bold=True, size=14, color="0077B5")
    row += 1
    
    sorted_results = sorted(results, key=lambda x: x.get('match_score', 0), reverse=True)
    
    medals = ["🥇", "🥈", "🥉", "4️⃣", "5️⃣"]
    for idx, result in enumerate(sorted_results[:5], 0):
        medal = medals[idx] if idx < len(medals) else f"{idx+1}."
        candidate_summary = f"{medal} {result.get('full_name', 'Unknown')} - {result.get('match_score', 0):.1f}%"
        ws[f'A{row}'] = candidate_summary
        ws[f'A{row}'].font = Font(bold=True, size=11)
        row += 1
        
        details = f"      {result.get('headline', 'N/A')}"
        if result.get('current_company'):
            details += f" | Current: {result.get('current_company')}"
        details += f" | Skills: {result.get('skills_score', 0):.0f}% | "
        details += f"Experience: {result.get('experience_score', 0):.0f}%"
        
        ws[f'A{row}'] = details
        ws[f'A{row}'].font = Font(size=9, italic=True)
        row += 1
        
        # Key strengths
        if result.get('key_strengths'):
            ws[f'A{row}'] = f"      Strengths: {result['key_strengths'][:200]}..."
            ws[f'A{row}'].font = Font(size=9, color="006100")
            row += 1
        
        row += 1
    
    row += 1
    
    # Job Description snippet
    if job_description:
        ws[f'A{row}'] = "📋 JOB DESCRIPTION OVERVIEW"
        ws[f'A{row}'].font = Font(bold=True, size=14, color="0077B5")
        row += 1
        
        jd_snippet = job_description[:500] + "..." if len(job_description) > 500 else job_description
        ws[f'A{row}'] = jd_snippet
        ws[f'A{row}'].font = Font(size=9, italic=True)
        ws[f'A{row}'].alignment = Alignment(wrap_text=True, vertical='top')
        row += 2


def create_linkedin_rankings_sheet(wb: Workbook, results: List[Dict[str, Any]]) -> None:
    """Create candidate rankings sheet"""
    
    ws = wb.create_sheet("Candidate Rankings")
    
    header_fill = PatternFill(start_color="0077B5", end_color="0077B5", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    
    headers = [
        'Rank', 'Full Name', 'Overall Match (%)', 'Skills Score (%)', 
        'Experience Score (%)', 'Education Score (%)', 'LinkedIn URL',
        'Headline', 'Current Company', 'Location', 'Recommendation'
    ]
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    sorted_results = sorted(results, key=lambda x: x.get('match_score', 0), reverse=True)
    
    for row_idx, result in enumerate(sorted_results, 2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=2, value=result.get('full_name', 'Unknown'))
        ws.cell(row=row_idx, column=3, value=round(result.get('match_score', 0), 1))
        ws.cell(row=row_idx, column=4, value=round(result.get('skills_score', 0), 1))
        ws.cell(row=row_idx, column=5, value=round(result.get('experience_score', 0), 1))
        ws.cell(row=row_idx, column=6, value=round(result.get('education_score', 0), 1))
        ws.cell(row=row_idx, column=7, value=result.get('profile_url', ''))
        ws.cell(row=row_idx, column=8, value=result.get('headline', ''))
        ws.cell(row=row_idx, column=9, value=result.get('current_company', ''))
        ws.cell(row=row_idx, column=10, value=result.get('location', ''))
        ws.cell(row=row_idx, column=11, value=result.get('recommendation', 'N/A'))
    
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx == 7:  # LinkedIn URL
            ws.column_dimensions[col_letter].width = 40
        elif col_idx == 8:  # Headline
            ws.column_dimensions[col_letter].width = 30
        else:
            ws.column_dimensions[col_letter].width = 18


def create_linkedin_skills_sheet(wb: Workbook, results: List[Dict[str, Any]]) -> None:
    """Create skills analysis sheet"""
    
    ws = wb.create_sheet("Skills Analysis")
    
    header_fill = PatternFill(start_color="0077B5", end_color="0077B5", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    
    headers = ['Rank', 'Full Name', 'Skills Score (%)', 'Key Strengths', 'Gaps / Missing Skills']
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    sorted_results = sorted(results, key=lambda x: x.get('match_score', 0), reverse=True)
    
    for row_idx, result in enumerate(sorted_results, 2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=2, value=result.get('full_name', 'Unknown'))
        ws.cell(row=row_idx, column=3, value=round(result.get('skills_score', 0), 1))
        
        strengths = result.get('key_strengths', 'N/A')
        ws.cell(row=row_idx, column=4, value=strengths)
        
        gaps = result.get('gaps', 'None identified')
        ws.cell(row=row_idx, column=5, value=gaps)
    
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 50
    ws.column_dimensions['E'].width = 50


def create_job_hopping_sheet(wb: Workbook, results: List[Dict[str, Any]]) -> None:
    """Create job hopping analysis sheet"""
    
    ws = wb.create_sheet("Job Hopping Analysis")
    
    header_fill = PatternFill(start_color="0077B5", end_color="0077B5", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    
    # Risk level color fills
    high_risk_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    medium_risk_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    low_risk_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    
    headers = ['Rank', 'Full Name', 'Job Hopping Risk', 'Risk Details', 'Recommendation']
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    sorted_results = sorted(results, key=lambda x: x.get('match_score', 0), reverse=True)
    
    for row_idx, result in enumerate(sorted_results, 2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=2, value=result.get('full_name', 'Unknown'))
        
        job_hopping_risk = result.get('job_hopping_risk', 'Unknown')
        risk_cell = ws.cell(row=row_idx, column=3, value=job_hopping_risk)
        risk_cell.font = Font(bold=True)
        
        # Apply color coding
        if job_hopping_risk and 'high' in job_hopping_risk.lower():
            risk_cell.fill = high_risk_fill
        elif job_hopping_risk and 'medium' in job_hopping_risk.lower():
            risk_cell.fill = medium_risk_fill
        elif job_hopping_risk and 'low' in job_hopping_risk.lower():
            risk_cell.fill = low_risk_fill
        
        detail = result.get('job_hopping_detail', 'No details available')
        detail_cell = ws.cell(row=row_idx, column=4, value=detail)
        detail_cell.alignment = Alignment(wrap_text=True, vertical='top')
        
        ws.cell(row=row_idx, column=5, value=result.get('recommendation', 'N/A'))
    
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 60
    ws.column_dimensions['E'].width = 20


def create_experience_details_sheet(wb: Workbook, results: List[Dict[str, Any]]) -> None:
    """Create experience details sheet"""
    
    ws = wb.create_sheet("Experience Details")
    
    header_fill = PatternFill(start_color="0077B5", end_color="0077B5", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    
    headers = [
        'Rank', 'Full Name', 'Experience Score (%)', 'Match Reasoning',
        'Current Company', 'Headline', 'Education Score (%)'
    ]
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    sorted_results = sorted(results, key=lambda x: x.get('match_score', 0), reverse=True)
    
    for row_idx, result in enumerate(sorted_results, 2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=2, value=result.get('full_name', 'Unknown'))
        ws.cell(row=row_idx, column=3, value=round(result.get('experience_score', 0), 1))
        
        reasoning = result.get('match_reasoning', 'No reasoning provided')
        reasoning_cell = ws.cell(row=row_idx, column=4, value=reasoning)
        reasoning_cell.alignment = Alignment(wrap_text=True, vertical='top')
        
        ws.cell(row=row_idx, column=5, value=result.get('current_company', 'N/A'))
        ws.cell(row=row_idx, column=6, value=result.get('headline', 'N/A'))
        ws.cell(row=row_idx, column=7, value=round(result.get('education_score', 0), 1))
    
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 60
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 30
    ws.column_dimensions['G'].width = 18


# ==============================================================================
# MODULE INITIALIZATION
# ==============================================================================

print("\n" + "="*80)
print("✅ LINKEDIN PROCESSING MODULE LOADED")
print("="*80)
print("\n📚 Available Functions:")
print("  • export_linkedin_analysis_to_excel() - Export to Excel with 5 sheets")
print("\n💡 Features:")
print("  ✓ Summary page with analysis overview")
print("  ✓ Candidate rankings with scores")
print("  ✓ Skills analysis")
print("  ✓ Job hopping analysis")
print("  ✓ Experience details")
print("\n🚀 Ready to export LinkedIn scout results!")
print("="*80 + "\n")
