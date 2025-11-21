# worker/sentiment_processing.py
"""
Complete sentiment processing with candidate-aware analysis and rate limit handling.
"""

import os
import sys
import io
import json
import uuid
import time
import random
import traceback
from datetime import datetime, date
from statistics import mean
from typing import List, Dict
from collections import defaultdict

# ensure repo root imports work
REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

import config
from utils import vertex_ai_utils, file_utils

# Google clients
from google.cloud import storage, bigquery
from google.api_core import exceptions as google_exceptions

# Initialize clients
STORAGE = storage.Client(project=config.PROJECT_ID)
BQ = bigquery.Client(project=config.PROJECT_ID)

# Local fallback dir
LOCAL_REPORT_DIR = os.path.abspath(os.path.join(REPO_ROOT, "local_reports"))
os.makedirs(LOCAL_REPORT_DIR, exist_ok=True)

# ---------- BigQuery Table Schema Initialization ----------

def ensure_bigquery_tables():
    """Ensure all required BigQuery tables exist with proper schemas."""
    dataset_id = f"{config.PROJECT_ID}.{config.BQ_DATASET}"
    
    try:
        BQ.get_dataset(dataset_id)
        print(f"✅ Dataset {dataset_id} exists")
    except google_exceptions.NotFound:
        print(f"⚠️ Dataset {dataset_id} not found. Creating...")
        dataset = bigquery.Dataset(dataset_id)
        dataset.location = config.LOCATION
        BQ.create_dataset(dataset, timeout=30)
        print(f"✅ Created dataset {dataset_id}")
    
    # Define schemas for ALL THREE tables
    schemas = {
        "sentiment_jobs": [
            bigquery.SchemaField("job_id", "STRING", mode="REQUIRED"),
            bigquery.SchemaField("ingestion_id", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("user_email", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("organization_id", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("status", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("created_at", "TIMESTAMP", mode="NULLABLE"),
            bigquery.SchemaField("updated_at", "TIMESTAMP", mode="NULLABLE"),
            bigquery.SchemaField("record_date", "DATE", mode="NULLABLE"),
            bigquery.SchemaField("record_date_str", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("report_url", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("error_message", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("meta", "STRING", mode="NULLABLE"),
        ],
        "sentiment_results": [
            bigquery.SchemaField("result_id", "STRING", mode="REQUIRED"),
            bigquery.SchemaField("job_id", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("ingestion_id", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("processed_by", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("created_at", "TIMESTAMP", mode="NULLABLE"),
            bigquery.SchemaField("record_date", "DATE", mode="NULLABLE"),
            bigquery.SchemaField("record_date_str", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("organization_id", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("user_email", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("unit_level", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("questionnaire_id", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("questionnaire_name", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("total_texts", "INTEGER", mode="NULLABLE"),
            bigquery.SchemaField("positive_count", "INTEGER", mode="NULLABLE"),
            bigquery.SchemaField("neutral_count", "INTEGER", mode="NULLABLE"),
            bigquery.SchemaField("negative_count", "INTEGER", mode="NULLABLE"),
            bigquery.SchemaField("average_confidence", "FLOAT", mode="NULLABLE"),
            bigquery.SchemaField("sentiment_index", "FLOAT", mode="NULLABLE"),
            bigquery.SchemaField("details", "RECORD", mode="REPEATED", fields=[
                bigquery.SchemaField("text_id", "STRING", mode="NULLABLE"),
                bigquery.SchemaField("question_id", "STRING", mode="NULLABLE"),
                bigquery.SchemaField("question_text", "STRING", mode="NULLABLE"),
                bigquery.SchemaField("text_snippet", "STRING", mode="NULLABLE"),
                bigquery.SchemaField("label", "STRING", mode="NULLABLE"),
                bigquery.SchemaField("confidence", "FLOAT", mode="NULLABLE"),
                bigquery.SchemaField("candidate_id", "STRING", mode="NULLABLE"),
                bigquery.SchemaField("candidate_name", "STRING", mode="NULLABLE"),
            ]),
            bigquery.SchemaField("meta", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("summary_text", "STRING", mode="NULLABLE"),  # ✅ FIX #1: You caught this!
        ],
        # ✅ FIX #2: Complete journal_vectors schema
        "journal_vectors": [
            bigquery.SchemaField("pdf_name", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("title", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("embedding", "FLOAT", mode="REPEATED"),  # Vector embeddings for RAG
            bigquery.SchemaField("summary", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("source_type", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("created_at", "TIMESTAMP", mode="NULLABLE"),
            bigquery.SchemaField("file_hash", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("organization_id", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("user_email", "STRING", mode="NULLABLE"),
        ]
    }
    
    for table_name, schema in schemas.items():
        table_id = f"{dataset_id}.{table_name}"
        try:
            BQ.get_table(table_id)
            print(f"✅ Table {table_name} exists")
        except google_exceptions.NotFound:
            print(f"⚠️ Table {table_name} not found. Creating...")
            table = bigquery.Table(table_id, schema=schema)
            BQ.create_table(table)
            print(f"✅ Created table {table_name}")

# ---------- BigQuery helpers ----------

def bq_insert_with_retry(BQ_client, table: str, rows: list, max_retries: int = 3, delay: float = 1.0):
    """Insert rows using insert_rows_json, with retries and explicit error logging."""
    if not rows:
        print("⚠️ bq_insert_with_retry: No rows to insert")
        return True
    
    print(f"DEBUG_BQ: Attempting to insert {len(rows)} row(s) to {table}")
    
    last_err = None
    for attempt in range(1, max_retries + 1):
        try:
            errors = BQ_client.insert_rows_json(table, rows)
            
            if not errors:
                print(f"✅ BQ insert SUCCESS on attempt {attempt}/{max_retries} for table {table}")
                return True
            else:
                print(f"❌ BQ insert returned errors on attempt {attempt}/{max_retries}:")
                for error in errors:
                    print(f"   Error: {error}")
                last_err = errors
                
            time.sleep(delay * (2 ** (attempt - 1)))
            
        except Exception as e:
            last_err = e
            print(f"❌ BQ insert EXCEPTION on attempt {attempt}/{max_retries}: {type(e).__name__}: {e}")
            traceback.print_exc()
            time.sleep(delay * (2 ** (attempt - 1)))
    
    error_msg = f"BQ insert FAILED for table {table} after {max_retries} attempts. Last error: {last_err}"
    print(f"🚨 {error_msg}")
    raise RuntimeError(error_msg)

def bq_query_with_retry(BQ_client, sql: str, job_config=None, max_retries: int = 3, delay: float = 1.0):
    """Run a query with retries, logging exceptions."""
    print(f"DEBUG_BQ: Executing query:\n{sql[:500]}")
    
    last_exc = None
    for attempt in range(1, max_retries + 1):
        try:
            q = BQ_client.query(sql, job_config=job_config)
            res = q.result()
            print(f"✅ BQ query SUCCESS on attempt {attempt}/{max_retries}")
            return res
        except Exception as e:
            last_exc = e
            print(f"❌ BQ query EXCEPTION on attempt {attempt}/{max_retries}: {type(e).__name__}: {e}")
            traceback.print_exc()
            time.sleep(delay * (2 ** (attempt - 1)))
    
    error_msg = f"BQ query FAILED after {max_retries} attempts. Last error: {last_exc}"
    print(f"🚨 {error_msg}")
    raise RuntimeError(error_msg)

# ---------- Email helper ----------
def _send_email_with_link(to_email: str, url: str, subject: str, body_html: str):
    """Send email (or fake it in local mode)"""
    if os.environ.get("LOCAL_EMAIL", "true").lower() in ("true", "1", "yes"):
        print(f"[LOCAL EMAIL] to={to_email} subject={subject}\n  link={url}\n  body={body_html[:200]}")
        return True
    raise NotImplementedError("Email sending not implemented for non-local mode")

# ---------- Sanitize text ----------
def _sanitize_text(text: str, include_full_text: bool = False):
    """Sanitize and redact PII from text"""
    try:
        redacted, flags = file_utils.sanitize_text(text, redact_names=True)
        if include_full_text:
            return text, flags
        else:
            return redacted, flags
    except Exception:
        import re
        red = re.sub(r'[\w\.-]+@[\w\.-]+', '[email_redacted]', text)
        if include_full_text:
            return text, {}
        return (red[:400], {})

# ---------- NEW: Candidate-level aggregation ----------

def aggregate_candidate_sentiment(details: List[Dict], domain: str = "exit_feedback") -> List[Dict]:
    """
    Aggregate per-question sentiments into per-candidate overall sentiment.
    """
    print(f"\n📊 Aggregating sentiment by candidate...")
    
    # Group by candidate
    candidate_groups = defaultdict(list)
    for detail in details:
        cid = detail.get("candidate_id", "unknown")
        candidate_groups[cid].append(detail)
    
    print(f"   Found {len(candidate_groups)} unique candidates")
    
    candidate_summaries = []
    
    for candidate_id, responses in candidate_groups.items():
        if not responses:
            continue
        
        candidate_name = responses[0].get("candidate_name", "Unknown")
        
        total = len(responses)
        pos = sum(1 for r in responses if r.get("label") == "positive")
        neg = sum(1 for r in responses if r.get("label") == "negative")
        neu = sum(1 for r in responses if r.get("label") == "neutral")
        
        confidences = [r.get("confidence", 0.5) for r in responses]
        avg_conf = mean(confidences) if confidences else 0.5
        
        sentiment_score = (pos - neg) / total if total > 0 else 0
        
        if sentiment_score > 0.3:
            overall_label = "positive"
        elif sentiment_score < -0.3:
            overall_label = "negative"
        else:
            overall_label = "neutral"
        
        question_breakdown = [
            {
                "question_id": r.get("question_id"),
                "question_text": r.get("question_text", ""),
                "response_snippet": r.get("text_snippet", "")[:200],
                "sentiment": r.get("label"),
                "confidence": r.get("confidence")
            }
            for r in responses
        ]
        
        candidate_summaries.append({
            "candidate_id": candidate_id,
            "candidate_name": candidate_name,
            "total_responses": total,
            "positive_count": pos,
            "negative_count": neg,
            "neutral_count": neu,
            "positive_pct": (pos / total * 100) if total > 0 else 0,
            "negative_pct": (neg / total * 100) if total > 0 else 0,
            "neutral_pct": (neu / total * 100) if total > 0 else 0,
            "overall_sentiment": overall_label,
            "overall_confidence": float(avg_conf),
            "sentiment_index": float(sentiment_score),
            "question_breakdown": question_breakdown
        })
    
    candidate_summaries.sort(key=lambda x: x["sentiment_index"])
    
    print(f"✅ Aggregated {len(candidate_summaries)} candidate summaries")
    return candidate_summaries


# ---------- Excel Report Builder ----------

def build_candidate_level_excel(candidate_summaries: List[Dict], 
                                 overall_details: List[Dict],
                                 job_id: str,
                                 filename: str,
                                 domain: str,
                                 output_path: str) -> str:
    """Build comprehensive Excel with CANDIDATE-LEVEL analysis."""
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill, Font
    
    print(f"\n📊 Building candidate-level Excel report...")
    
    total_candidates = len(candidate_summaries)
    total_responses = sum(c["total_responses"] for c in candidate_summaries)
    avg_responses = total_responses / total_candidates if total_candidates > 0 else 0
    pos_candidates = sum(1 for c in candidate_summaries if c["overall_sentiment"] == "positive")
    neg_candidates = sum(1 for c in candidate_summaries if c["overall_sentiment"] == "negative")
    neu_candidates = sum(1 for c in candidate_summaries if c["overall_sentiment"] == "neutral")
    
    # Sheet 1: Executive Summary
    exec_data = {
        "Metric": [
            "File Name", "Processing Date", "Domain", "Job ID", "─────────────────",
            "Total Candidates", "Total Responses", "Avg Responses/Candidate", "─────────────────",
            "Positive Candidates", "Negative Candidates", "Neutral Candidates", "─────────────────",
            "Positive %", "Negative %", "Neutral %", "─────────────────",
            "High Risk Candidates (>60% negative)", "Medium Risk Candidates (40-60% negative)"
        ],
        "Value": [
            filename, datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"), domain, job_id, "",
            total_candidates, total_responses, f"{avg_responses:.1f}", "",
            pos_candidates, neg_candidates, neu_candidates, "",
            f"{(pos_candidates/total_candidates*100):.1f}%",
            f"{(neg_candidates/total_candidates*100):.1f}%",
            f"{(neu_candidates/total_candidates*100):.1f}%", "",
            sum(1 for c in candidate_summaries if c["negative_pct"] > 60),
            sum(1 for c in candidate_summaries if 40 < c["negative_pct"] <= 60)
        ]
    }
    df_exec = pd.DataFrame(exec_data)
    
    # Sheet 2: Candidate Summary
    candidate_data = []
    for c in candidate_summaries:
        if c["negative_pct"] > 60:
            risk_flag = "⚠️ HIGH"
        elif c["negative_pct"] > 40:
            risk_flag = "⚡ MEDIUM"
        else:
            risk_flag = "✓ LOW"
        
        candidate_data.append({
            "Candidate ID": c["candidate_id"],
            "Candidate Name": c["candidate_name"],
            "Total Responses": c["total_responses"],
            "Overall Sentiment": c["overall_sentiment"].upper(),
            "Sentiment Index": f"{c['sentiment_index']:.3f}",
            "Positive": c["positive_count"],
            "Negative": c["negative_count"],
            "Neutral": c["neutral_count"],
            "Positive %": f"{c['positive_pct']:.1f}%",
            "Negative %": f"{c['negative_pct']:.1f}%",
            "Avg Confidence": f"{c['overall_confidence']:.3f}",
            "Risk Flag": risk_flag
        })
    df_candidates = pd.DataFrame(candidate_data)
    
    # Sheet 3: Detailed Feedback
    detailed_data = []
    for c in candidate_summaries:
        for q in c["question_breakdown"]:
            detailed_data.append({
                "Candidate ID": c["candidate_id"],
                "Candidate Name": c["candidate_name"],
                "Question": q["question_text"],
                "Response": q["response_snippet"],
                "Sentiment": q["sentiment"].upper(),
                "Confidence": f"{q['confidence']:.3f}"
            })
    df_detailed = pd.DataFrame(detailed_data)
    
    # Sheet 4: Question Analysis
    question_stats = defaultdict(lambda: {"pos": 0, "neg": 0, "neu": 0, "total": 0})
    for c in candidate_summaries:
        for q in c["question_breakdown"]:
            qid = q["question_id"]
            question_stats[qid]["total"] += 1
            if q["sentiment"] == "positive":
                question_stats[qid]["pos"] += 1
            elif q["sentiment"] == "negative":
                question_stats[qid]["neg"] += 1
            else:
                question_stats[qid]["neu"] += 1
    
    question_data = []
    for qid, stats in question_stats.items():
        total = stats["total"]
        pos_pct = (stats["pos"]/total*100) if total > 0 else 0
        neg_pct = (stats["neg"]/total*100) if total > 0 else 0
        
        if pos_pct > 60:
            trend = "👍 Strong Positive"
        elif neg_pct > 40:
            trend = "👎 Concerning"
        else:
            trend = "➖ Mixed"
        
        question_data.append({
            "Question": qid,
            "Total Responses": total,
            "Positive": stats["pos"],
            "Negative": stats["neg"],
            "Neutral": stats["neu"],
            "Positive %": f"{pos_pct:.1f}%",
            "Negative %": f"{neg_pct:.1f}%",
            "Sentiment Trend": trend
        })
    df_questions = pd.DataFrame(question_data)
    
    # Sheet 5: Risk Alerts
    risk_data = []
    for c in candidate_summaries:
        if c["negative_pct"] > 40:
            neg_questions = [q for q in c["question_breakdown"] if q["sentiment"] == "negative"]
            top_concerns = ", ".join([q["question_text"][:30] for q in neg_questions[:3]])
            
            risk_data.append({
                "Candidate ID": c["candidate_id"],
                "Candidate Name": c["candidate_name"],
                "Risk Level": "HIGH" if c["negative_pct"] > 60 else "MEDIUM",
                "Negative %": f"{c['negative_pct']:.1f}%",
                "Sentiment Index": f"{c['sentiment_index']:.3f}",
                "Negative Responses": c["negative_count"],
                "Key Concerns": top_concerns if top_concerns else "Multiple concerns"
            })
    
    if risk_data:
        df_risk = pd.DataFrame(risk_data)
    else:
        df_risk = pd.DataFrame([{"Message": "✅ No high-risk candidates identified"}])
    
    # Write all sheets
    print(f"   Writing Excel sheets...")
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_exec.to_excel(writer, sheet_name="Executive Summary", index=False)
        df_candidates.to_excel(writer, sheet_name="Candidate Summary", index=False)
        df_detailed.to_excel(writer, sheet_name="Detailed Feedback", index=False)
        df_questions.to_excel(writer, sheet_name="Question Analysis", index=False)
        df_risk.to_excel(writer, sheet_name="Risk Alerts", index=False)
    
    # Apply formatting
    print(f"   Applying conditional formatting...")
    wb = load_workbook(output_path)
    
    if "Candidate Summary" in wb.sheetnames:
        ws = wb["Candidate Summary"]
        for row in range(2, len(candidate_data) + 2):
            sentiment_cell = ws.cell(row, 4)
            if sentiment_cell.value == "POSITIVE":
                sentiment_cell.fill = PatternFill(start_color="90EE90", fill_type="solid")
            elif sentiment_cell.value == "NEGATIVE":
                sentiment_cell.fill = PatternFill(start_color="FFB6C1", fill_type="solid")
            
            risk_cell = ws.cell(row, 12)
            if "HIGH" in str(risk_cell.value):
                risk_cell.fill = PatternFill(start_color="FF6B6B", fill_type="solid")
                risk_cell.font = Font(bold=True, color="FFFFFF")
            elif "MEDIUM" in str(risk_cell.value):
                risk_cell.fill = PatternFill(start_color="FFD93D", fill_type="solid")
                risk_cell.font = Font(bold=True)
    
    if "Risk Alerts" in wb.sheetnames and risk_data:
        ws = wb["Risk Alerts"]
        for row in range(2, len(risk_data) + 2):
            risk_cell = ws.cell(row, 3)
            if risk_cell.value == "HIGH":
                for col in range(1, 8):
                    ws.cell(row, col).fill = PatternFill(start_color="FFE5E5", fill_type="solid")
    
    # Auto-adjust column widths
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 100)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(output_path)
    print(f"✅ Candidate-level Excel report saved: {output_path}")
    print(f"   📊 Report contains 5 sheets:")
    print(f"      1. Executive Summary - Key metrics")
    print(f"      2. Candidate Summary - {len(candidate_data)} candidates")
    print(f"      3. Detailed Feedback - {len(detailed_data)} responses")
    print(f"      4. Question Analysis - {len(question_data)} questions")
    print(f"      5. Risk Alerts - {len(risk_data)} at-risk candidates")
    
    return output_path


# ---------- MAIN PROCESSING FUNCTION ----------

def process_job(payload: dict, sync_mode: bool = False):
    """Process sentiment analysis job with CANDIDATE-LEVEL aggregation."""
    print(f"\n{'='*80}")
    print(f"🚀 START process_job | job_id={payload.get('job_id')} | sync_mode={sync_mode}")
    print(f"{'='*80}\n")
    
    job_id = payload.get("job_id") or str(uuid.uuid4())
    filename = payload.get("filename", "upload.xlsx")
    domain = payload.get("domain", "general")
    user_email = payload.get("user_email", "guest_user@gmail.com")
    org_id = payload.get("organization_id", "guest_org")
    include_full_text = payload.get("include_full_text", False)

    now = datetime.utcnow()
    today = date.today()
    record_date = today.isoformat()
    record_date_str = today.strftime("%d-%m-%Y")

    # STEP 0: Ensure BigQuery tables exist
    try:
        print("🔧 Checking/creating BigQuery tables...")
        ensure_bigquery_tables()
    except Exception as e:
        error_msg = f"Failed to initialize BigQuery tables: {e}"
        print(f"🚨 {error_msg}")
        traceback.print_exc()
        return None, f"error: {error_msg}"

    # STEP 1: Insert job row to BigQuery
    job_row = {
        "job_id": job_id,
        "ingestion_id": payload.get("ingestion_id", str(uuid.uuid4())),
        "user_email": user_email,
        "organization_id": org_id,
        "status": "PROCESSING",
        "created_at": now.isoformat(),
        "updated_at": now.isoformat(),
        "record_date": record_date,
        "record_date_str": record_date_str,
        "meta": json.dumps({"filename": filename, "domain": domain})
    }

    table_job = f"{config.PROJECT_ID}.{config.BQ_DATASET}.sentiment_jobs"
    
    try:
        print(f"📤 Inserting job row to {table_job}...")
        bq_insert_with_retry(BQ, table_job, [job_row])
        print("✅ Job row inserted successfully")
    except Exception as e:
        error_msg = f"CRITICAL: Failed to insert job row to BigQuery: {e}"
        print(f"🚨 {error_msg}")
        traceback.print_exc()
        return None, f"error: {error_msg}"

    # STEP 2: Parse uploaded file
    try:
        print(f"📂 Parsing file: {filename}...")
        file_bytes = payload.get("file_bytes")
        if isinstance(file_bytes, str):
            file_bytes = bytes.fromhex(file_bytes)
        sheets = file_utils.read_excel_sheets_from_bytes(file_bytes, filename)
        print(f"✅ Parsed {len(sheets)} sheet(s)")
    except Exception as e:
        error_msg = f"Failed to parse file: {e}"
        print(f"❌ {error_msg}")
        traceback.print_exc()
        
        return None, f"error: {error_msg}"

    # STEP 3: Process sheets WITH CANDIDATE AWARENESS AND BATCH PROCESSING
    overall_details = []
    all_candidate_summaries = []
    total_all = 0
    pos_all = 0
    neu_all = 0
    neg_all = 0
    confidences_all = []

    try:
        for sheet_idx, (sheet_name, df) in enumerate(sheets.items(), 1):
            print(f"\n📊 Processing sheet {sheet_idx}/{len(sheets)}: {sheet_name}")
            
            # 🆕 USE CANDIDATE-AWARE EXTRACTION
            rows = file_utils.extract_candidate_feedback_rows(df)
            print(f"   Found {len(rows)} feedback rows from candidates")
            
            # Track unique candidates in this sheet
            unique_candidates = set(r.get("candidate_id") for r in rows)
            print(f"   Unique candidates in sheet: {len(unique_candidates)}")
            
            # 🆕 BATCH PROCESS ALL ROWS AT ONCE (not one by one!)
            details = vertex_ai_utils.process_feedback_batch(rows, domain, include_full_text)
            
            # Count totals from the batch results
            total_all += len(details)
            for d in details:
                if d["label"] == "positive":
                    pos_all += 1
                elif d["label"] == "negative":
                    neg_all += 1
                else:
                    neu_all += 1
                
                try:
                    confidences_all.append(float(d["confidence"]))
                except Exception:
                    pass

            # 🆕 AGGREGATE BY CANDIDATE FOR THIS SHEET
            sheet_candidate_summaries = aggregate_candidate_sentiment(details, domain)
            all_candidate_summaries.extend(sheet_candidate_summaries)
            
            # Questionnaire/Sheet-level summary
            total_q = len(details)
            pos_q = sum(1 for d in details if d["label"] == "positive")
            neu_q = sum(1 for d in details if d["label"] == "neutral")
            neg_q = sum(1 for d in details if d["label"] == "negative")
            avg_conf_q = mean([d["confidence"] for d in details]) if details else 0.0
            sentiment_index_q = (pos_q - neg_q) / total_q if total_q > 0 else 0.0

            questionnaire_summary = {
                "result_id": str(uuid.uuid4()),
                "job_id": job_id,
                "ingestion_id": job_row["ingestion_id"],
                "processed_by": "sentiment_agent_v2_candidate_aware",
                "created_at": now.isoformat(),
                "record_date": record_date,
                "record_date_str": record_date_str,
                "organization_id": org_id,
                "user_email": user_email,
                "unit_level": "questionnaire",
                "questionnaire_id": str(uuid.uuid4())[:12],
                "questionnaire_name": sheet_name,
                "total_texts": total_q,
                "positive_count": pos_q,
                "neutral_count": neu_q,
                "negative_count": neg_q,
                "average_confidence": float(avg_conf_q),
                "sentiment_index": float(sentiment_index_q),
                "details": details[:100],  # Only first 100 for BQ limits
                "meta": json.dumps({"domain": domain, "unique_candidates": len(unique_candidates)})
            }

            # Insert questionnaire summary
            table_results = f"{config.PROJECT_ID}.{config.BQ_DATASET}.sentiment_results"
            try:
                print(f"📤 Inserting questionnaire summary to BigQuery...")
                bq_insert_with_retry(BQ, table_results, [questionnaire_summary])
                print(f"✅ Questionnaire summary inserted")
            except Exception as e:
                print(f"🚨 FAILED to insert questionnaire summary: {e}")
                raise

            overall_details.extend(details)

        # 🆕 OVERALL CANDIDATE AGGREGATION (across all sheets)
        print(f"\n📊 Computing overall candidate aggregation across all sheets...")
        final_candidate_summaries = aggregate_candidate_sentiment(overall_details, domain)
        
        # Overall summary
        total_overall = total_all
        pos_overall = pos_all
        neu_overall = neu_all
        neg_overall = neg_all
        avg_conf_overall = mean(confidences_all) if confidences_all else 0.0
        sentiment_index_overall = (pos_overall - neg_overall) / total_overall if total_overall > 0 else 0.0

        overall_result = {
            "result_id": str(uuid.uuid4()),
            "job_id": job_id,
            "ingestion_id": job_row["ingestion_id"],
            "processed_by": "sentiment_agent_v2_candidate_aware",
            "created_at": now.isoformat(),
            "record_date": record_date,
            "record_date_str": record_date_str,
            "organization_id": org_id,
            "user_email": user_email,
            "unit_level": "doc",
            "questionnaire_id": None,
            "questionnaire_name": filename,
            "total_texts": total_overall,
            "positive_count": pos_overall,
            "neutral_count": neu_overall,
            "negative_count": neg_overall,
            "average_confidence": float(avg_conf_overall),
            "sentiment_index": float(sentiment_index_overall),
            "details": overall_details[:100],
            "meta": json.dumps({"domain": domain, "total_candidates": len(final_candidate_summaries)})
        }
        
        print(f"\n📤 Inserting overall result to BigQuery...")
        bq_insert_with_retry(BQ, table_results, [overall_result])
        print(f"✅ Overall result inserted")
        
        # Insert aggregated learning summary to journal_vectors
        try:
            print(f"\n🧠 Recording learning summary to journal_vectors...")
            
            sentiment_distribution = {
                "positive_pct": (pos_overall / total_overall * 100) if total_overall > 0 else 0,
                "negative_pct": (neg_overall / total_overall * 100) if total_overall > 0 else 0,
                "neutral_pct": (neu_overall / total_overall * 100) if total_overall > 0 else 0
            }
            
            # Detect domains
            print(f"   🔍 Detecting feedback domains...")
            detected_domains = set()
            domain_keywords = {
                "training_feedback": ["training", "trainer", "course", "workshop", "learning"],
                "onboarding_feedback": ["onboarding", "first day", "orientation", "laptop", "access"],
                "exit_feedback": ["leaving", "exit", "resignation", "quit"],
                "manager_feedback": ["manager", "supervisor", "boss", "leadership"],
                "appraisal_feedback": ["performance", "appraisal", "review", "evaluation"],
                "culture_feedback": ["culture", "work environment", "team", "collaboration"],
                "interview_feedback": ["interview", "candidate", "hiring", "recruitment"],
                "pulse_feedback": ["survey", "engagement", "satisfaction", "morale"]
            }
            
            for detail in overall_details:
                text_lower = detail.get("text_snippet", "").lower()
                for domain_type, keywords in domain_keywords.items():
                    if any(keyword in text_lower for keyword in keywords):
                        detected_domains.add(domain_type)
            
            if not detected_domains:
                detected_domains.add(domain)
            
            detected_domains_list = sorted(list(detected_domains))
            print(f"   ✅ Detected domains: {', '.join(detected_domains_list)}")
            
            # Generate AI summary with retry
            print(f"   🤖 Generating AI insights summary...")
            
            high_risk_candidates = [c for c in final_candidate_summaries if c["negative_pct"] > 60][:3]
            positive_candidates = [c for c in final_candidate_summaries if c["overall_sentiment"] == "positive"][:3]
            
            ai_summary_prompt = f"""
            Analyze the following candidate sentiment analysis results and provide a concise 3-4 sentence summary.
            
            Context:
            - Total Candidates: {len(final_candidate_summaries)}
            - Total Responses: {total_overall}
            - Positive: {pos_overall} ({sentiment_distribution['positive_pct']:.1f}%), Negative: {neg_overall} ({sentiment_distribution['negative_pct']:.1f}%), Neutral: {neu_overall} ({sentiment_distribution['neutral_pct']:.1f}%)
            - Domains: {', '.join(detected_domains_list)}
            - Sentiment Index: {sentiment_index_overall:.2f}
            - High Risk Candidates: {len(high_risk_candidates)}
            
            High Risk Candidates: {', '.join([f"{c['candidate_name']} ({c['negative_pct']:.0f}% negative)" for c in high_risk_candidates])}
            Positive Candidates: {', '.join([f"{c['candidate_name']}" for c in positive_candidates])}
            
            Write in business-appropriate style starting with "Overall..." capturing candidate-level trends, risk areas, and actionable insights.
            
            Summary:
            """
            
            # Retry logic for AI summary
            max_retries = 3
            ai_summary = None
            for attempt in range(max_retries):
                try:
                    if attempt > 0:
                        wait_time = (2 ** attempt) + random.random()
                        print(f"   ⏳ Rate limit hit, waiting {wait_time:.1f}s before retry {attempt + 1}/{max_retries}...")
                        time.sleep(wait_time)
                    
                    ai_summary = vertex_ai_utils.generate_freeform(ai_summary_prompt)
                    
                    if "Gemini freeform failed" in ai_summary or "429" in ai_summary:
                        if attempt < max_retries - 1:
                            continue
                        else:
                            ai_summary = None
                    else:
                        ai_summary = ai_summary.strip()
                        print(f"   ✅ AI summary generated ({len(ai_summary)} chars)")
                        break
                        
                except Exception as e:
                    print(f"   ⚠️ AI summary attempt {attempt + 1} failed: {str(e)[:100]}")
                    if attempt == max_retries - 1:
                        ai_summary = None
            
            # Fallback summary
            if not ai_summary or len(ai_summary) < 50:
                print(f"   ℹ️ Using fallback summary")
                
                if sentiment_index_overall > 0.3:
                    ai_summary = f"Overall positive candidate sentiment with {pos_overall} positive responses ({sentiment_distribution['positive_pct']:.0f}%) across {len(final_candidate_summaries)} candidates. "
                    ai_summary += f"{len(high_risk_candidates)} candidates require attention due to high negative feedback."
                elif sentiment_index_overall < -0.3:
                    ai_summary = f"Overall negative candidate sentiment with {neg_overall} negative responses ({sentiment_distribution['negative_pct']:.0f}%). "
                    ai_summary += f"{len(high_risk_candidates)} high-risk candidates identified requiring immediate follow-up."
                else:
                    ai_summary = f"Mixed candidate sentiment across {len(final_candidate_summaries)} candidates with {pos_overall} positive, {neg_overall} negative responses. "
                    ai_summary += f"{len(high_risk_candidates)} candidates flagged for attention."
            
            # Create learning summary
            learning_summary = {
                "ai_generated_summary": ai_summary,
                "detected_domains": detected_domains_list,
                "primary_domain": domain,
                "total_candidates": len(final_candidate_summaries),
                "high_risk_candidates": len(high_risk_candidates),
                "job_summary": f"Processed {total_overall} responses from {len(final_candidate_summaries)} candidates",
                "sentiment_index": sentiment_index_overall,
                "distribution": sentiment_distribution,
                "avg_overall_confidence": avg_conf_overall,
                "filename": filename,
                "job_id": job_id
            }
            
            # Generate embedding
            import hashlib
            content_hash = hashlib.md5(f"{job_id}_{filename}_{now.isoformat()}".encode()).hexdigest()[:12]
            
            embedding_text = f"""
            Candidate Sentiment Analysis Summary
            {ai_summary}
            
            Domain: {', '.join(detected_domains_list)}
            File: {filename}
            Candidates: {len(final_candidate_summaries)}
            Total Responses: {total_overall}
            Positive: {pos_overall} ({sentiment_distribution['positive_pct']:.1f}%), Negative: {neg_overall} ({sentiment_distribution['negative_pct']:.1f}%)
            Sentiment Index: {sentiment_index_overall:.3f}
            High Risk: {len(high_risk_candidates)} candidates
            Organization: {org_id}
            """.strip()
            
            print(f"   📢 Generating embedding vector...")
            embedding_vector = vertex_ai_utils.generate_embedding(embedding_text)
            
            if not embedding_vector:
                print(f"   ⚠️ Embedding generation failed, using empty vector")
                embedding_vector = []
            else:
                print(f"   ✅ Generated embedding vector (dim: {len(embedding_vector)})")
            
            # Insert into journal_vectors
            journal_row = {
                "pdf_name": f"sentiment_summary_{content_hash}",
                "title": f"Candidate Sentiment Analysis: {filename} [{', '.join(detected_domains_list)}]",
                "embedding": embedding_vector,
                "summary": json.dumps(learning_summary),
                "source_type": "sentiment_analysis",
                "created_at": now.isoformat(),
                "file_hash": content_hash,
                "organization_id": org_id,
                "user_email": user_email
            }
            
            journal_table = f"{config.PROJECT_ID}.{config.BQ_DATASET}.journal_vectors"
            bq_insert_with_retry(BQ, journal_table, [journal_row])
            print(f"✅ Learning summary recorded to journal_vectors")
            
        except Exception as e:
            print(f"⚠️ Failed to insert learning summary: {str(e)[:200]}")
            traceback.print_exc()

    except Exception as e:
        error_msg = f"Processing/BigQuery error: {e}"
        print(f"🚨 {error_msg}")
        traceback.print_exc()
        
        return None, f"error: {error_msg}"

    # STEP 4: Build comprehensive CANDIDATE-LEVEL Excel report
    try:
        print(f"\n📊 Building candidate-level Excel report...")
        
        local_path = os.path.join(LOCAL_REPORT_DIR, f"sentiment-candidate-report-{job_id}.xlsx")
        
        build_candidate_level_excel(
            candidate_summaries=final_candidate_summaries,
            overall_details=overall_details,
            job_id=job_id,
            filename=filename,
            domain=domain,
            output_path=local_path
        )
        
        report_url = f"file://{local_path}"
        print(f"✅ Report URL: {report_url}")
        
    except Exception as e:
        print(f"⚠️ Report generation failed: {e}")
        traceback.print_exc()
        local_path = os.path.join(LOCAL_REPORT_DIR, f"sentiment-candidate-report-{job_id}.xlsx")
        report_url = f"file://{local_path}"

    # STEP 5: Send email
    try:
        high_risk_count = sum(1 for c in final_candidate_summaries if c["negative_pct"] > 60)
        
        subject = f"Candidate Sentiment Analysis Complete - Job {job_id}"
        body_html = f"""
        <h2>Candidate Sentiment Analysis Results</h2>
        <p><strong>Total Candidates:</strong> {len(final_candidate_summaries)}</p>
        <p><strong>Total Responses:</strong> {total_overall}</p>
        <hr>
        <p><strong>Overall Sentiment Distribution:</strong></p>
        <p>✅ Positive: {pos_overall} ({pos_overall/total_overall*100:.1f}%)</p>
        <p>➖ Neutral: {neu_overall} ({neu_overall/total_overall*100:.1f}%)</p>
        <p>❌ Negative: {neg_overall} ({neg_overall/total_overall*100:.1f}%)</p>
        <hr>
        <p><strong>Risk Assessment:</strong></p>
        <p>⚠️ High Risk Candidates: {high_risk_count}</p>
        <p><strong>Sentiment Index:</strong> {sentiment_index_overall:.3f}</p>
        <hr>
        <p><a href="{report_url}">Download Detailed Report</a></p>
        <p><em>Report includes candidate-level analysis, question breakdown, and risk alerts.</em></p>
        """
        _send_email_with_link(user_email, report_url, subject, body_html)
    except Exception as e:
        print(f"⚠️ Email send failed: {e}")

    # STEP 6: Job completion
    try:
        print(f"\n📤 Recording job completion...")
        print("✅ Job completed (status update skipped due to streaming buffer)")
    except Exception as e:
        print(f"⚠️ Note: {e}")

    # Final return
    summary_text = f"Processed {len(final_candidate_summaries)} candidates with {total_overall} responses: +{pos_overall}/-{neg_overall}/~{neu_overall} (index: {sentiment_index_overall:.3f}). {high_risk_count} high-risk candidates identified."
    
    print(f"\n{'='*80}")
    print(f"✅ COMPLETED process_job | result_id={overall_result.get('result_id')}")
    print(f"   Summary: {summary_text}")
    print(f"   📊 Candidate Breakdown:")
    print(f"      - Total Candidates: {len(final_candidate_summaries)}")
    print(f"      - Positive Sentiment: {sum(1 for c in final_candidate_summaries if c['overall_sentiment'] == 'positive')}")
    print(f"      - Negative Sentiment: {sum(1 for c in final_candidate_summaries if c['overall_sentiment'] == 'negative')}")
    print(f"      - Neutral Sentiment: {sum(1 for c in final_candidate_summaries if c['overall_sentiment'] == 'neutral')}")
    print(f"      - High Risk (>60% neg): {high_risk_count}")
    print(f"{'='*80}\n")
    
    return overall_result.get("result_id"), summary_text