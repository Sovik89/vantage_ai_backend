import hashlib
from io import BytesIO
from docx import Document
import pypdf
import PyPDF2
import pandas as pd
import os
import traceback
import io
import re
from google.cloud import storage
import spacy
from typing import List, Tuple

nlp = None
def _ensure_spacy():
    global nlp
    if nlp is None:
        try:
            nlp = spacy.load("en_core_web_sm")
        except Exception:
            # fallback: load blank English pipeline if model not present
            nlp = spacy.blank("en")
    return nlp

# This file is now fully synchronous and works with in-memory BytesIO buffers.

def compute_file_hash(content: bytes) -> str:
    """Computes a SHA256 hash of the file's content."""
    return hashlib.sha256(content).hexdigest()

def get_text_from_file(content_buffer: BytesIO, file_ext: str) -> str:
    """Routes a file buffer to the correct text extraction function."""
    file_ext = file_ext.lower().replace('.', '')
    if file_ext == "pdf":
        return read_pdf_text(content_buffer)
    elif file_ext == "docx":
        return read_docx_text(content_buffer)
    elif file_ext == "txt":
        return read_txt_text(content_buffer)
    elif file_ext == "csv":
        return read_csv_text(content_buffer)
    else:
        print(f"⚠️ Unsupported file type provided.")
        return ""

def read_pdf_text(content_buffer: BytesIO) -> str:
    """Extracts text from a PDF using a primary and fallback library."""
    text = ""
    try:
        content_buffer.seek(0)
        pdf_reader = pypdf.PdfReader(content_buffer)
        text = "".join(page.extract_text() or "" for page in pdf_reader.pages)
    except Exception as e_pypdf:
        print(f"Primary PDF reader (pypdf) failed: {e_pypdf}. Trying fallback.")
        try:
            content_buffer.seek(0)
            pdf_reader_fallback = PyPDF2.PdfReader(content_buffer)
            text = "".join(page.extract_text() or "" for page in pdf_reader_fallback.pages)
        except Exception as e_pypdf2:
            print(f"Fallback PDF reader (PyPDF2) also failed: {e_pypdf2}")
            return ""
    return text.strip()

def read_docx_text(content_buffer: BytesIO) -> str:
    """Extracts text from a DOCX file."""
    try:
        content_buffer.seek(0)
        doc = Document(content_buffer)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        return text.strip()
    except Exception as e:
        print(f"Could not extract text from DOCX: {e}")
        return ""

def read_txt_text(content_buffer: BytesIO) -> str:
    """Reads text from a TXT file."""
    try:
        content_buffer.seek(0)
        content = content_buffer.read()
        return content.decode("utf-8").strip()
    except Exception as e:
        print(f"Could not read TXT: {e}")
        return ""

def read_csv_text(content_buffer: BytesIO) -> str:
    """Reads the first 20 rows of a CSV."""
    try:
        content_buffer.seek(0)
        df = pd.read_csv(content_buffer)
        return df.head(20).to_markdown(index=False)
    except Exception as e:
        print(f"Could not read CSV: {e}")
        return ""

# Note: The export_to_docx function remains unchanged as it doesn't handle async file objects.
def export_to_docx(records, consolidated_text=None, user_prompt=None):
    """Exports summaries and a consolidated answer to a DOCX file in memory."""
    doc = Document()
    doc.add_heading("Vantage.Ai Insights", 0)
    
    if user_prompt:
        doc.add_paragraph(f"User Query: {user_prompt}")
    
    if consolidated_text:
        doc.add_heading("Consolidated Summary", level=1)
        doc.add_paragraph(consolidated_text)
    
    if records:
        doc.add_heading("Source Document Summaries", level=1)
        for r in records:
            title = r.get("title") or r.get("pdf_name") or "Untitled"
            doc.add_heading(title, level=2)
            if r.get("summary"):
                doc.add_paragraph(r["summary"])

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def clean_text(text):
    if not text:
        return ""
    return (
        text.replace("–", "-")
            .replace("—", "-")
            .replace("’", "'")
            .replace("“", '"')
            .replace("”", '"')
            .replace("…", "...")
    )

def export_to_pdf(records, consolidated_text=None, user_prompt=None):
    from fpdf import FPDF
    from io import BytesIO

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, clean_text("HR Journal Insights"), ln=True, align="C")

    pdf.set_font("Arial", "", 12)
    if user_prompt:
        pdf.ln(10)
        pdf.multi_cell(0, 10, clean_text(f"User Query: {user_prompt}"))

    if consolidated_text:
        pdf.ln(10)
        pdf.set_font("Arial", "B", 14)
        pdf.cell(0, 10, clean_text("Consolidated Summary"), ln=True)
        pdf.set_font("Arial", "", 12)
        pdf.multi_cell(0, 10, clean_text(consolidated_text))

    pdf.ln(10)
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, clean_text("Individual Paper Summaries"), ln=True)

    for r in records:
        title = r.get("title") or r.get("pdf_name") or "Untitled"
        summary = r.get("summary", "")
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 10, clean_text(title), ln=True)
        pdf.set_font("Arial", "", 12)
        if summary:
            pdf.multi_cell(0, 10, clean_text(summary))
        pdf.ln(5)

    buf = BytesIO()
    buf.write(pdf.output(dest='S').encode('latin-1'))
    buf.seek(0)
    return buf, None, True

def extract_text_from_bytes(file_bytes: bytes, filename: str, max_rows:int=10000) -> Tuple[List[str], List[int]]:
    """
    Returns (texts, row_indices) - detects a 'feedback' column or picks first string column.
    """
    buf = io.BytesIO(file_bytes)
    try:
        df = pd.read_excel(buf, engine="openpyxl")
    except Exception:
        df = pd.read_excel(buf)  # fallback
    return _extract_from_df(df, max_rows)

def extract_text_from_gcs(gcs_path: str, max_rows:int=10000) -> Tuple[List[str], List[int]]:
    # gcs_path like gs://bucket/path
    client = storage.Client()
    if not gcs_path.startswith("gs://"):
        raise ValueError("gcs_path must be gs://")
    _p = gcs_path[len("gs://"):]
    bucket_name, _, blob_path = _p.partition("/")
    bucket = client.bucket(bucket_name)
    blob = bucket.blob(blob_path)
    data = blob.download_as_bytes()
    return extract_text_from_bytes(data, blob_path, max_rows)

def _extract_from_df(df: pd.DataFrame, max_rows:int):
    # detect feedback-like column
    candidates = [c for c in df.columns if any(k in c.lower() for k in ("feedback","comment","response","text","answer"))]
    if candidates:
        col = candidates[0]
        series = df[col]
    else:
        # fallback: pick first object/string column
        obj_cols = [c for c in df.columns if df[c].dtype == 'object']
        if obj_cols:
            col = obj_cols[0]
            series = df[col]
        else:
            # if no text columns, stringify rows
            series = df.apply(lambda r: " | ".join([str(x) for x in r.values if pd.notna(x)]), axis=1)
    # clean and limit
    texts = []
    indices = []
    for i, v in enumerate(series.fillna("").astype(str).tolist(), start=1):
        t = v.strip()
        if not t:
            continue
        if len(texts) >= max_rows:
            break
        texts.append(t)
        indices.append(i)
    return texts, indices

def sanitize_text(text:str, redact_names:bool=True) -> Tuple[str, dict]:
    """
    Redact PERSON and EMAIL found in text.
    Returns (redacted_text, flags)
    """
    nlp = _ensure_spacy()
    doc = nlp(text)
    redacted = text
    flags = {"redacted_names": False, "redacted_emails": False}
    # replace entities in reverse order to keep indices valid
    for ent in reversed(list(doc.ents)):
        if ent.label_ in ("PERSON",):
            if redact_names:
                start, end = ent.start_char, ent.end_char
                redacted = redacted[:start] + "[REDACTED_NAME]" + redacted[end:]
                flags["redacted_names"] = True
    # redact emails with regex as extra layer
    email_re = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")
    if email_re.search(redacted):
        redacted = email_re.sub("[REDACTED_EMAIL]", redacted)
        flags["redacted_emails"] = True
    return redacted, flags